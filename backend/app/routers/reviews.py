from collections import Counter
from datetime import date, timedelta
from io import BytesIO
import json
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select

from app.ai_client import ai_config, call_chat_completion
from app.audit import add_audit_log
from app.config import get_settings
from app.database import get_session
from app.deps import get_current_user
from app.models import (
    DailyPlan,
    DiaryEntry,
    ImprovementAction,
    PlanStep,
    ReviewReport,
    Task,
    TaskLog,
    User,
    WeaknessRecord,
    utc_now,
)
from app.schemas import (
    ReviewGenerateRequest,
    ReviewRead,
    ReviewSaveResponse,
    ReviewSmartSummaryResponse,
    ReviewUpdate,
)

router = APIRouter(prefix="/api/reviews", tags=["reviews"])


def today() -> date:
    return date.today()


def resolve_range(payload: ReviewGenerateRequest) -> tuple[str, str]:
    current = today()
    if payload.start_date and payload.end_date:
        return payload.start_date, payload.end_date
    if payload.report_type == "daily":
        value = current.isoformat()
        return value, value
    if payload.report_type == "weekly":
        start = current - timedelta(days=current.weekday())
        return start.isoformat(), current.isoformat()
    if payload.report_type == "monthly":
        return current.replace(day=1).isoformat(), current.isoformat()
    return current.isoformat(), current.isoformat()


def get_report_for_user(session: Session, report_id: int, user_id: int) -> ReviewReport:
    report = session.get(ReviewReport, report_id)
    if report is None or report.user_id != user_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Review not found")
    return report


def build_summary(
    session: Session,
    *,
    user_id: int,
    start_date: str,
    end_date: str,
) -> dict[str, object]:
    weaknesses = session.exec(
        select(WeaknessRecord).where(
            WeaknessRecord.user_id == user_id,
            WeaknessRecord.is_deleted == False,  # noqa: E712
            WeaknessRecord.occur_date >= start_date,
            WeaknessRecord.occur_date <= end_date,
        )
    ).all()
    weakness_ids = [item.id for item in weaknesses if item.id is not None]
    actions = session.exec(
        select(ImprovementAction).where(ImprovementAction.user_id == user_id)
    ).all()
    scoped_actions = [action for action in actions if action.weakness_id in weakness_ids]
    completed_actions = [action for action in scoped_actions if action.status == "已完成"]
    unfinished_actions = [action for action in scoped_actions if action.status != "已完成"]

    plans = session.exec(
        select(DailyPlan).where(
            DailyPlan.user_id == user_id,
            DailyPlan.plan_date >= start_date,
            DailyPlan.plan_date <= end_date,
        )
    ).all()
    plan_ids = [plan.id for plan in plans if plan.id is not None]
    plan_steps = (
        session.exec(
            select(PlanStep).where(
                PlanStep.user_id == user_id,
                PlanStep.plan_id.in_(plan_ids),
            )
        ).all()
        if plan_ids
        else []
    )

    tasks = session.exec(
        select(Task).where(Task.user_id == user_id, Task.is_active == True)  # noqa: E712
    ).all()
    task_logs = session.exec(
        select(TaskLog).where(
            TaskLog.user_id == user_id,
            TaskLog.log_date >= start_date,
            TaskLog.log_date <= end_date,
        )
    ).all()
    diaries = session.exec(
        select(DiaryEntry).where(
            DiaryEntry.user_id == user_id,
            DiaryEntry.is_deleted == False,  # noqa: E712
            DiaryEntry.entry_date >= start_date,
            DiaryEntry.entry_date <= end_date,
        )
    ).all()

    category_counts = Counter(item.category for item in weaknesses)
    top_category = category_counts.most_common(1)[0][0] if category_counts else "暂无"
    highest = max(weaknesses, key=lambda item: item.severity, default=None)
    completed_steps = sum(1 for step in plan_steps if step.status == "completed")
    checked_task_ids = {log.task_id for log in task_logs if not log.is_makeup}

    return {
        "start_date": start_date,
        "end_date": end_date,
        "weakness_count": len(weaknesses),
        "top_category": top_category,
        "category_counts": dict(category_counts),
        "highest_severity": {
            "title": highest.title,
            "severity": highest.severity,
            "category": highest.category,
        }
        if highest
        else None,
        "unfinished_actions": [
            {"id": action.id, "title": action.title, "status": action.status}
            for action in unfinished_actions
        ],
        "completed_actions": [
            {"id": action.id, "title": action.title}
            for action in completed_actions
        ],
        "planning": {
            "plans": len(plans),
            "steps": len(plan_steps),
            "completed_steps": completed_steps,
        },
        "tasks": {
            "active_tasks": len(tasks),
            "checked_tasks": len(checked_task_ids),
            "checkin_logs": len(task_logs),
        },
        "diary_count": len(diaries),
        "weaknesses": [
            {
                "title": item.title,
                "category": item.category,
                "severity": item.severity,
                "occur_date": item.occur_date,
                "behavior": item.behavior,
                "reason": item.reason,
                "correction": item.correction,
            }
            for item in weaknesses
        ],
    }


def summary_to_text(summary: dict[str, object]) -> str:
    planning = summary["planning"]
    tasks = summary["tasks"]
    highest = summary["highest_severity"]
    lines = [
        f"时间范围：{summary['start_date']} 至 {summary['end_date']}",
        f"缺点记录数量：{summary['weakness_count']}",
        f"最高频问题类型：{summary['top_category']}",
        "最高严重程度问题："
        + (
            f"{highest['title']}（{highest['category']}，严重程度 {highest['severity']}）"
            if highest
            else "暂无"
        ),
        f"未完成改正措施：{len(summary['unfinished_actions'])}",
        f"已完成改正措施：{len(summary['completed_actions'])}",
        (
            f"规划完成情况：{planning['plans']} 个规划，"
            f"{planning['completed_steps']}/{planning['steps']} 个步骤已完成"
        ),
        (
            f"任务完成情况：{tasks['checked_tasks']}/{tasks['active_tasks']} 个活跃任务有打卡，"
            f"共 {tasks['checkin_logs']} 条打卡记录"
        ),
        f"日记数量：{summary['diary_count']}",
    ]
    return "\n".join(lines)


def build_prompt(report_type: str, summary: dict[str, object]) -> str:
    period_label = {"daily": "今日", "weekly": "本周", "monthly": "本月"}.get(report_type, "本阶段")
    weaknesses = summary["weaknesses"]
    weakness_lines = "\n".join(
        f"- {item['occur_date']}｜{item['category']}｜严重度 {item['severity']}｜{item['title']}：{item['behavior']}"
        for item in weaknesses
    ) or "- 暂无缺点记录"
    unfinished = "\n".join(f"- {item['title']}（{item['status']}）" for item in summary["unfinished_actions"]) or "- 暂无"
    completed = "\n".join(f"- {item['title']}" for item in summary["completed_actions"]) or "- 暂无"

    return (
        f"你是一位严谨、具体、重视执行的个人复盘教练。请根据以下本地记录，为我完成一份{period_label}复盘。\n\n"
        f"时间范围：{summary['start_date']} 至 {summary['end_date']}\n\n"
        f"本地摘要：\n{summary_to_text(summary)}\n\n"
        f"缺点记录：\n{weakness_lines}\n\n"
        f"高频问题：{summary['top_category']}\n\n"
        f"已采取措施：\n{completed}\n\n"
        f"未完成措施：\n{unfinished}\n\n"
        "请用中文输出以下内容：\n"
        "1. 问题总结：按重要程度列出 3-5 个核心问题。\n"
        "2. 根本原因：不要泛泛而谈，要指出触发场景、认知偏差和行为链条。\n"
        "3. 可执行改正措施：每条措施要具体到时间、动作和验收标准。\n"
        "4. 行动清单：如果是今日复盘，输出明日行动清单；如果是周/月复盘，输出下周或下月行动清单。\n"
        "5. 风险提醒：指出最可能复发的问题和预防办法。\n"
        "请保持直接、务实，不要输出鸡汤。"
    )


def create_review_report(
    session: Session,
    *,
    user_id: int,
    report_type: str,
    start_date: str,
    end_date: str,
    input_summary: str,
    ai_prompt: str | None = None,
    ai_response: str | None = None,
    accepted: bool = False,
) -> ReviewReport:
    now = utc_now()
    report = ReviewReport(
        user_id=user_id,
        report_type=report_type,
        start_date=start_date,
        end_date=end_date,
        input_summary=input_summary,
        ai_prompt=ai_prompt,
        ai_response=ai_response,
        final_content=ai_response,
        accepted=accepted,
        next_actions=None,
        created_at=now,
        updated_at=now,
    )
    session.add(report)
    session.flush()
    return report


def call_ai_summary(prompt: str) -> str:
    settings = get_settings()
    return call_chat_completion(
        ai_config(settings),
        [
            {"role": "system", "content": "你是严谨、具体、重视执行的中文复盘助手。"},
            {"role": "user", "content": prompt},
        ],
        temperature=0.3,
        timeout=45.0,
        failure_message="AI 总结失败，请检查配置后重试",
    )


def latest_report_for_range(
    session: Session,
    *,
    user_id: int,
    report_type: str,
    start_date: str,
    end_date: str,
) -> ReviewReport | None:
    return session.exec(
        select(ReviewReport)
        .where(
            ReviewReport.user_id == user_id,
            ReviewReport.report_type == report_type,
            ReviewReport.start_date == start_date,
            ReviewReport.end_date == end_date,
        )
        .order_by(ReviewReport.created_at.desc())
    ).first()


def build_review_docx(summary: dict[str, object], report: ReviewReport | None = None) -> BytesIO:
    from docx import Document

    planning = summary["planning"]
    tasks = summary["tasks"]
    weaknesses = summary["weaknesses"]
    unfinished_actions = summary["unfinished_actions"]
    completed_actions = summary["completed_actions"]
    title = "初记复盘文档"

    document = Document()
    document.add_heading(title, 0)
    document.add_heading("时间范围", level=1)
    document.add_paragraph(f"{summary['start_date']} 至 {summary['end_date']}")

    document.add_heading("今日/本周完成情况", level=1)
    document.add_paragraph(
        f"规划：{planning['plans']} 个规划，{planning['completed_steps']}/{planning['steps']} 个安排已完成。"
    )
    document.add_paragraph(
        f"任务：{tasks['checked_tasks']}/{tasks['active_tasks']} 个活跃任务有打卡，共 {tasks['checkin_logs']} 条记录。"
    )
    document.add_paragraph(f"日记：{summary['diary_count']} 篇。")

    document.add_heading("缺点与问题", level=1)
    if weaknesses:
        for item in weaknesses:
            document.add_paragraph(
                f"{item['occur_date']}｜{item['category']}｜严重度 {item['severity']}｜{item['title']}：{item['behavior']}",
                style="List Bullet",
            )
            if item["reason"]:
                document.add_paragraph(f"原因：{item['reason']}")
    else:
        document.add_paragraph("暂无缺点与问题记录。")

    document.add_heading("改正措施", level=1)
    if completed_actions:
        document.add_paragraph("已完成：")
        for item in completed_actions:
            document.add_paragraph(item["title"], style="List Bullet")
    if unfinished_actions:
        document.add_paragraph("未完成：")
        for item in unfinished_actions:
            document.add_paragraph(f"{item['title']}（{item['status']}）", style="List Bullet")
    if not completed_actions and not unfinished_actions:
        document.add_paragraph("暂无改正措施。")

    document.add_heading("后续计划", level=1)
    if report and report.final_content:
        document.add_paragraph(report.final_content)
    elif unfinished_actions:
        document.add_paragraph("优先完成以下未完成措施：")
        for item in unfinished_actions:
            document.add_paragraph(item["title"], style="List Bullet")
    else:
        document.add_paragraph("保持当前节奏，下一阶段继续记录、复盘并补充可执行安排。")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("", response_model=list[ReviewRead])
def list_reviews(
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> list[ReviewReport]:
    return list(
        session.exec(
            select(ReviewReport)
            .where(ReviewReport.user_id == current_user.id)
            .order_by(ReviewReport.created_at.desc())
        ).all()
    )


@router.get("/{review_id}", response_model=ReviewRead)
def get_review(
    review_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ReviewReport:
    return get_report_for_user(session, review_id, current_user.id)


@router.put("/{review_id}", response_model=ReviewRead)
def update_review(
    review_id: int,
    payload: ReviewUpdate,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ReviewReport:
    report = get_report_for_user(session, review_id, current_user.id)
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(report, field, value)
    report.updated_at = utc_now()
    session.add(report)
    add_audit_log(
        session,
        request=request,
        action="review.update",
        user_id=current_user.id,
        detail={"report_id": report.id},
    )
    session.commit()
    session.refresh(report)
    return report


@router.post("/generate-local-summary")
def generate_local_summary(
    payload: ReviewGenerateRequest,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> dict[str, object]:
    start_date, end_date = resolve_range(payload)
    summary = build_summary(session, user_id=current_user.id, start_date=start_date, end_date=end_date)
    return {
        "report_type": payload.report_type,
        "start_date": start_date,
        "end_date": end_date,
        "summary": summary,
        "input_summary": summary_to_text(summary),
    }


@router.post("/generate-ai-prompt")
def generate_ai_prompt(
    payload: ReviewGenerateRequest,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> dict[str, object]:
    start_date, end_date = resolve_range(payload)
    summary = build_summary(session, user_id=current_user.id, start_date=start_date, end_date=end_date)
    input_summary = summary_to_text(summary)
    prompt = build_prompt(payload.report_type, summary)
    now = utc_now()
    report = ReviewReport(
        user_id=current_user.id,
        report_type=payload.report_type,
        start_date=start_date,
        end_date=end_date,
        input_summary=input_summary,
        ai_prompt=prompt,
        ai_response=None,
        final_content=None,
        accepted=False,
        next_actions=None,
        created_at=now,
        updated_at=now,
    )
    session.add(report)
    session.flush()
    add_audit_log(
        session,
        request=request,
        action="review.generate_prompt",
        user_id=current_user.id,
        detail={"report_id": report.id, "report_type": report.report_type},
    )
    session.commit()
    session.refresh(report)
    return {
        "report_id": report.id,
        "report_type": report.report_type,
        "start_date": report.start_date,
        "end_date": report.end_date,
        "input_summary": input_summary,
        "ai_prompt": prompt,
        "summary": summary,
    }


@router.post("/generate-smart-summary", response_model=ReviewSmartSummaryResponse)
def generate_smart_summary(
    payload: ReviewGenerateRequest,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> dict[str, object]:
    start_date, end_date = resolve_range(payload)
    summary = build_summary(session, user_id=current_user.id, start_date=start_date, end_date=end_date)
    input_summary = summary_to_text(summary)
    summary_material = build_prompt(payload.report_type, summary)
    settings = get_settings()

    if ai_config(settings).configured:
        ai_summary = call_ai_summary(summary_material)
        report = create_review_report(
            session,
            user_id=current_user.id,
            report_type=payload.report_type,
            start_date=start_date,
            end_date=end_date,
            input_summary=input_summary,
            ai_prompt=summary_material,
            ai_response=ai_summary,
            accepted=True,
        )
        message = "智能总结已生成并保存"
        ai_enabled = True
    else:
        report = create_review_report(
            session,
            user_id=current_user.id,
            report_type=payload.report_type,
            start_date=start_date,
            end_date=end_date,
            input_summary=input_summary,
            ai_prompt=summary_material,
        )
        ai_summary = None
        message = "未配置智能总结密钥，已生成可复制的总结素材"
        ai_enabled = False

    add_audit_log(
        session,
        request=request,
        action="review.smart_summary",
        user_id=current_user.id,
        detail={"report_id": report.id, "report_type": report.report_type, "ai_enabled": ai_enabled},
    )
    session.commit()
    session.refresh(report)
    return {
        "report_id": report.id,
        "report_type": report.report_type,
        "start_date": report.start_date,
        "end_date": report.end_date,
        "input_summary": input_summary,
        "ai_summary": ai_summary,
        "summary_material": None if ai_enabled else summary_material,
        "ai_enabled": ai_enabled,
        "message": message,
    }


@router.post("/document")
def generate_review_document(
    payload: ReviewGenerateRequest,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    start_date, end_date = resolve_range(payload)
    summary = build_summary(session, user_id=current_user.id, start_date=start_date, end_date=end_date)
    report = latest_report_for_range(
        session,
        user_id=current_user.id,
        report_type=payload.report_type,
        start_date=start_date,
        end_date=end_date,
    )
    buffer = build_review_docx(summary, report)
    filename = f"初记复盘文档-{start_date}-至-{end_date}.docx"
    add_audit_log(
        session,
        request=request,
        action="review.document",
        user_id=current_user.id,
        detail={"report_type": payload.report_type, "start_date": start_date, "end_date": end_date},
    )
    session.commit()
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/save-ai-response", response_model=ReviewRead)
def save_ai_response(
    payload: ReviewSaveResponse,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> ReviewReport:
    if payload.report_id is not None:
        report = get_report_for_user(session, payload.report_id, current_user.id)
    else:
        start_date = payload.start_date or today().isoformat()
        end_date = payload.end_date or start_date
        now = utc_now()
        report = ReviewReport(
            user_id=current_user.id,
            report_type=payload.report_type,
            start_date=start_date,
            end_date=end_date,
            input_summary=None,
            ai_prompt=payload.ai_prompt,
            created_at=now,
            updated_at=now,
        )
        session.add(report)
        session.flush()

    report.ai_prompt = payload.ai_prompt or report.ai_prompt
    report.ai_response = payload.ai_response
    report.final_content = payload.final_content or payload.ai_response
    report.accepted = payload.accepted
    report.next_actions = payload.next_actions
    report.updated_at = utc_now()
    session.add(report)
    add_audit_log(
        session,
        request=request,
        action="review.save_response",
        user_id=current_user.id,
        detail={"report_id": report.id, "accepted": report.accepted},
    )
    session.commit()
    session.refresh(report)
    return report


@router.delete("/{review_id}")
def delete_report(
    review_id: int,
    request: Request,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
) -> dict[str, bool]:
    report = get_report_for_user(session, review_id, current_user.id)
    session.delete(report)
    add_audit_log(
        session, request=request, action="review.delete",
        user_id=current_user.id, detail={"report_id": review_id},
    )
    session.commit()
    return {"ok": True}
